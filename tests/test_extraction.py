import os
import pytest
import docx
import pptx
from extraction.router import extract_file
from schemas import ExtractionResult, TextBlock
from extraction.retention import calculate_retention

@pytest.fixture
def sample_docx(tmp_path):
    path = tmp_path / "test.docx"
    doc = docx.Document()
    doc.add_paragraph("Hello world DOCX text.")
    table = doc.add_table(rows=1, cols=2)
    row = table.rows[0].cells
    row[0].text = "Cell 1"
    row[1].text = "Cell 2"
    doc.save(path)
    return str(path)

@pytest.fixture
def sample_pptx(tmp_path):
    path = tmp_path / "test.pptx"
    prs = pptx.Presentation()
    blank_slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_slide_layout)
    txBox = slide.shapes.add_textbox(0, 0, 100, 100)
    tf = txBox.text_frame
    tf.text = "Hello world PPTX text."
    prs.save(path)
    return str(path)

# Creating a PDF is harder without an external library, we'll test docx and pptx and retention.
# We will create a dummy PDF file that is invalid to test error handling at least.
@pytest.fixture
def invalid_pdf(tmp_path):
    path = tmp_path / "test.pdf"
    with open(path, "w") as f:
        f.write("Not a real pdf")
    return str(path)

def test_extract_docx(sample_docx):
    res = extract_file(sample_docx)
    assert res.success == True
    assert len(res.text_blocks) >= 2
    texts = [b.text for b in res.text_blocks]
    assert "Hello world DOCX text." in texts
    assert "Cell 1\\tCell 2" in texts or "Cell 1\tCell 2" in texts

def test_extract_pptx(sample_pptx):
    res = extract_file(sample_pptx)
    assert res.success == True
    assert len(res.text_blocks) >= 1
    texts = [b.text for b in res.text_blocks]
    assert any("Hello world PPTX text." in t for t in texts)

def test_extract_invalid_pdf(invalid_pdf):
    res = extract_file(invalid_pdf)
    assert res.success == False
    assert len(res.errors) > 0
    assert "Failed to open PDF" in res.errors[0]

def test_extract_not_found():
    res = extract_file("does_not_exist.pdf")
    assert res.success == False
    assert "File not found" in res.errors[0]

def test_retention():
    # Test text
    ref = "This is a test of text retention. It should match 100%!"
    ext = [
        TextBlock(id="1", file="f", block_type="text", source="n", text="This is a test of text retention."),
        TextBlock(id="2", file="f", block_type="text", source="n", text="It should match 100%!")
    ]
    ret_res = calculate_retention(ext, ref)
    assert ret_res["meets_target"] == True
    assert ret_res["retention_score"] == 100.0

def test_retention_partial():
    ref = "word1 word2 word3 word4 word5"
    ext = "word1 word2 word3"
    ret_res = calculate_retention(ext, ref)
    assert ret_res["meets_target"] == False
    assert ret_res["retention_score"] == 60.0
