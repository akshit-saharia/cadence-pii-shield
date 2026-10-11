import pptx
import uuid
from schemas import ExtractionResult, TextBlock, ExtractedImage

def extract_pptx(file_path: str) -> ExtractionResult:
    result = ExtractionResult(file=file_path, success=False)
    
    try:
        presentation = pptx.Presentation(file_path)
    except Exception as e:
        result.errors.append(f"Failed to open PPTX: {str(e)}")
        return result
        
    for slide_index, slide in enumerate(presentation.slides):
        slide_num = slide_index + 1
        has_text = False
        
        # Speaker notes
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes_text = slide.notes_slide.notes_text_frame.text.strip()
            if notes_text:
                result.text_blocks.append(TextBlock(
                    id=str(uuid.uuid4()),
                    file=file_path,
                    page_or_slide=slide_num,
                    text=notes_text,
                    source="native",
                    block_type="notes",
                ))
        
        # Shapes on slide
        for shape in slide.shapes:
            # Handle Group Shapes by recursively extracting them
            def extract_from_shape(shp, current_slide_num):
                nonlocal has_text
                
                if shp.shape_type == 6: # GROUP
                    for sub_shp in shp.shapes:
                        extract_from_shape(sub_shp, current_slide_num)
                        
                elif hasattr(shp, "text") and shp.text:
                    text = shp.text.strip()
                    if text:
                        has_text = True
                        bbox = None
                        if hasattr(shp, "left") and hasattr(shp, "top") and hasattr(shp, "width") and hasattr(shp, "height"):
                            try:
                                bbox = [shp.left, shp.top, shp.left + shp.width, shp.top + shp.height]
                            except:
                                pass
                                
                        result.text_blocks.append(TextBlock(
                            id=str(uuid.uuid4()),
                            file=file_path,
                            page_or_slide=current_slide_num,
                            text=text,
                            bbox=bbox,
                            source="native",
                            block_type="text",
                            metadata={"shape_name": shp.name}
                        ))
                        
                elif shp.has_table:
                    has_text = True
                    table_text = []
                    for row in shp.table.rows:
                        row_text = [cell.text_frame.text.strip() for cell in row.cells]
                        table_text.append("\t".join(row_text))
                    text = "\n".join(table_text)
                    if text.strip():
                        result.text_blocks.append(TextBlock(
                            id=str(uuid.uuid4()),
                            file=file_path,
                            page_or_slide=current_slide_num,
                            text=text,
                            source="native",
                            block_type="table",
                            metadata={"shape_name": shp.name}
                        ))
                        
                elif hasattr(shp, "image"):
                    try:
                        img_bytes = shp.image.blob
                        result.images.append(ExtractedImage(
                            id=str(uuid.uuid4()),
                            file=file_path,
                            page_or_slide=current_slide_num,
                            image_bytes=img_bytes,
                            metadata={"shape_name": shp.name, "ext": shp.image.ext}
                        ))
                    except Exception as e:
                        result.warnings.append(f"Failed to extract image from shape {shp.name} on slide {current_slide_num}: {str(e)}")

            extract_from_shape(shape, slide_num)
            
        if not has_text:
            result.warnings.append(f"Slide {slide_num} appears to have no native text or is image-only.")

    result.warnings.append("Note: Extraction of some SmartArt or complex embedded objects may not be fully supported.")
    
    result.success = True
    return result
