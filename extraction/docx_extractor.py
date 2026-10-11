import docx
import uuid
from schemas import ExtractionResult, TextBlock, ExtractedImage

def extract_docx(file_path: str) -> ExtractionResult:
    result = ExtractionResult(file=file_path, success=False)
    
    try:
        doc = docx.Document(file_path)
    except Exception as e:
        result.errors.append(f"Failed to open DOCX: {str(e)}")
        return result
        
    # Document properties
    try:
        prop = doc.core_properties
        result.metadata["title"] = prop.title
        result.metadata["author"] = prop.author
    except Exception as e:
        result.warnings.append(f"Could not read properties: {str(e)}")

    # Paragraphs
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if text:
            result.text_blocks.append(TextBlock(
                id=str(uuid.uuid4()),
                file=file_path,
                page_or_slide=None,
                text=text,
                source="native",
                block_type="text",
                metadata={"paragraph_index": i}
            ))
            
    # Tables
    for i, table in enumerate(doc.tables):
        table_text = []
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells]
            table_text.append("\t".join(row_text))
        
        text = "\n".join(table_text)
        if text.strip():
            result.text_blocks.append(TextBlock(
                id=str(uuid.uuid4()),
                file=file_path,
                page_or_slide=None,
                text=text,
                source="native",
                block_type="table",
                metadata={"table_index": i}
            ))

    # Headers and Footers
    for section in doc.sections:
        if section.header:
            for para in section.header.paragraphs:
                text = para.text.strip()
                if text:
                    result.text_blocks.append(TextBlock(
                        id=str(uuid.uuid4()),
                        file=file_path,
                        page_or_slide=None,
                        text=text,
                        source="native",
                        block_type="header"
                    ))
        if section.footer:
            for para in section.footer.paragraphs:
                text = para.text.strip()
                if text:
                    result.text_blocks.append(TextBlock(
                        id=str(uuid.uuid4()),
                        file=file_path,
                        page_or_slide=None,
                        text=text,
                        source="native",
                        block_type="footer"
                    ))

    # Extracted Images
    try:
        for rel in doc.part.rels.values():
            if "image" in rel.target_ref:
                img_bytes = rel.target_part.blob
                result.images.append(ExtractedImage(
                    id=str(uuid.uuid4()),
                    file=file_path,
                    image_bytes=img_bytes,
                    metadata={"target_ref": rel.target_ref}
                ))
    except Exception as e:
        result.warnings.append(f"Failed to extract images: {str(e)}")

    result.warnings.append("Note: Extraction of text boxes and floating objects is not fully supported by python-docx.")
    
    result.success = True
    return result
